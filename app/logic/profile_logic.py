from kivymd.uix.menu import MDDropdownMenu
from kivymd.uix.textfield import MDTextField
from kivymd.uix.dialog import MDDialog
from kivymd.uix.button import MDRaisedButton, MDFlatButton
from kivymd.uix.boxlayout import MDBoxLayout
from kivymd.toast import toast

from kivy.metrics import dp, sp
from kivy.lang import Builder
from kivy.factory import Factory # instances vers les .kv qui ont été chargés dans le main.py
from kivy.utils import platform
if platform == "android":
    from jnius import autoclass, cast
    from android import activity

from openpyxl import load_workbook
from plyer import filechooser
from collections import defaultdict
import os
import csv
from datetime import datetime

class ProfileController:
    """ Chargement des données et création des dictionnaires """
    
    def __init__(self, **kwargs):
        # Initialisation des variables
        self.profile_menu = None

        # --- Dictionnaire vide des perfs pour chaque exercice ---
        self.all_exercise_dict = dict()

    def on_profile_icon_press(self, button):
        """
        Ouvre un menu déroulant
        :param button: bouton cliqué
        :return:
        """
        if not self.profile_menu:
            menu_items = [
                {
                    "text": "Importer ton fichier",
                    "trailing_icon": "download",
                    "on_release": lambda: (
                        self.open_file_manager_import(),
                        self.profile_menu.dismiss()
                    )
                },
                {
                    "text": "Exporter ton fichier",
                    "trailing_icon": "file-plus",
                    "on_release": lambda: self.profile_menu_callback("Créé ton fichier")
                },
            ]

            self.profile_menu = MDDropdownMenu(
                caller=button,
                items=menu_items,
                width_mult=4,
                radius=dp(10),
            )

        self.profile_menu.open()

    def profile_menu_callback(self, text):
        print(f"Option sélectionnée : {text}")
        self.profile_menu.dismiss()
    
    # ---------- Import manuel du fichier ----------
    
    def open_file_manager_import(self):
        """Ouvre le sélecteur natif pour choisir un fichier Excel"""
        # Le sélecteur de fichier de plyer ne prend pas de chemin de départ sur Android
        # mais on peut filtrer les extensions
        if platform == "android" :
            self.open_file_android(self.select_file)
        else :
            filechooser.open_file(
                on_selection=self.select_file
            )

    def open_file_android(self, on_selection):
        PythonActivity = autoclass('org.kivy.android.PythonActivity')
        Intent = autoclass('android.content.Intent')

        intent = Intent(Intent.ACTION_GET_CONTENT)
        intent.setType(
            "*/*")  # tu peux restreindre à "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        intent.addCategory(Intent.CATEGORY_OPENABLE)

        # Stocke ton callback dans une variable globale
        global ANDROID_FILE_CALLBACK
        ANDROID_FILE_CALLBACK = on_selection

        # Attache un listener sur le résultat
        activity.bind(on_activity_result=self._on_activity_result)

        # Lance le chooser
        PythonActivity.mActivity.startActivityForResult(intent, 1001)

    def _on_activity_result(self, requestCode, resultCode, data):
        from android import activity
        if requestCode != 1001:
            return

        if resultCode == -1:  # RESULT_OK
            uri = data.getData().toString()
            ANDROID_FILE_CALLBACK([uri])
        else:
            ANDROID_FILE_CALLBACK([])

    def copy_uri_to_temp_file(self, uri_string):
        PythonActivity = autoclass('org.kivy.android.PythonActivity')
        context = PythonActivity.mActivity
        Uri = autoclass('android.net.Uri')
        uri = Uri.parse(uri_string)

        content_resolver = context.getContentResolver()
        input_stream = content_resolver.openInputStream(uri)

        temp_dir = context.getCacheDir().getAbsolutePath()
        temp_path = os.path.join(temp_dir, "imported_excel.xlsx")

        FileOutputStream = autoclass('java.io.FileOutputStream')
        output_stream = FileOutputStream(temp_path)

        buffer = bytearray(1024)
        read = input_stream.read(buffer)
        while read != -1:
            output_stream.write(buffer, 0, read)
            read = input_stream.read(buffer)

        input_stream.close()
        output_stream.close()

        return temp_path

    def select_file(self, selection):
        """Callback quand un fichier est sélectionné"""

        print(f"Sélection : {selection}")

        # selection est soit None, soit une liste vide, soit [chemin]
        if not selection:
            toast("Aucun fichier sélectionné")
            print("Aucun fichier sélectionné (ou annulation).")
            return

        file_path = selection[0]  # premier fichier choisi
        print("Fichier choisi :", file_path)
        if not file_path:
            toast("Fichier invalide")
            print("Fichier invalide")
            return

        # Si on est sur Android → convertir content:// en vrai fichier
        if file_path.startswith("content://"):
            print("URI détectée → copie vers fichier temporaire")
            file_path = self.copy_uri_to_temp_file(file_path)

        print("Chemin final utilisé :", file_path)

        # Maintenant safe : file_path est une string
        if any(file_path.endswith(ext) for ext in [".xls", ".xlsx", ".csv"]):
            # traiter le fichier
            try:
                # Charger le fichier
                self.all_exercise_dict = self.load_excel_as_dict(file_path)
                print(f"Dictionnaire complet importé : {self.all_exercise_dict}")

                database = self.convert_old_db(self.all_exercise_dict)
                # print(f"Dictionnaire converti : {database}")

                # # Récupérer le nom de la première feuille
                # first_key = list(self.all_exercise_dict.keys())[0]

                # # Ajouter dans l'onglet [VISUALISATION]
                # self.app.view.update_exercise_menu_button_text(first_key)
                # self.app.view.select_exercise(first_key)

                # Alerter l'utilisateur
                toast("Fichier chargé avec succès !")

                # Charger la base complète depuis le dossier CSV
                # database = self.import_database_from_csv(folder_path="my_training_data")

                # Vérifier la structure
                print(database.keys())
                print(f"Dictionnaire converti : {database}")
                
                # liste des noms des exercices
                exercise_names = self.get_exercise_names(database)
                print(f"Exercices dans la base : {exercise_names}")

                # Ajouter dans l'onglet [VISUALISATION]
                self.app.view.update_exercise_menu_button_text(exercise_names[0])
                self.app.view.select_exercise(exercise_names[0])

            except Exception as e:
                toast(f"Erreur : {str(e)}",
                      # background=[0.8, 0.3, 0.3, 1]
                      )
                print("erreur :" + str(e))
        else:
            toast(text="Extension non supportée")


    # Charger toutes les feuilles dans un dict[str, list[dict]]
    def load_excel_as_dict(self, path):  # -> dict[str, list[dict[str, str | float | datetime | None]]]:
        wb = load_workbook(path, data_only=True)
        # all_sheets: dict[str, list[dict[str, str | float | datetime | None]]] = {}
        all_sheets = {}

        for sheetname in wb.sheetnames:
            ws = wb[sheetname]

            # Récupérer les en-têtes (première ligne)
            headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]

            # Récupérer les lignes suivantes sous forme de dict
            rows = []
            for row in ws.iter_rows(min_row=2, values_only=True):
                row_dict = dict(zip(headers, row))
                rows.append(row_dict)

            all_sheets[sheetname] = rows

        return all_sheets

    def convert_old_db(self,all_exercise_dict):
        """
        Convertit l'ancienne base all_exercise_dict en une base moderne :
        
        database = {
            "exercise_library": {id: {name, muscle_group, primary_group, secondary_group, type}},
            "sessions": [
                {
                    "date": datetime,
                    "exercises": [
                        {
                            "exercise_id": id,
                            "sets": [{"r": reps, "w": weight}, ...]
                        }
                    ]
                }
            ]
        }
        """

        database = {
            "exercise_library": {},
            "sessions": [],
            "body_weight_history": []  # à compléter si tu as des données de poids corporel
        }

        # --- 1️⃣ Créer exercise_library ---
        exercise_ids = {}  # mapping name -> id
        for i, ex_name in enumerate(all_exercise_dict.keys(), start=1):
            exercise_ids[ex_name] = i

            # Ici tu peux personnaliser les infos sur chaque exercice
            # Pour l'instant on met des valeurs par défaut pour muscle_group, primary, secondary, type
            database["exercise_library"][i] = {
                "name": ex_name,
                "muscle_group": "Unknown",     # à compléter selon tes exercices
                "primary_group": [],           # muscles principaux
                "secondary_group": [],         # muscles secondaires
                "type": "bodyweight"           # ou "barbell" si connu
            }

        # --- 2️⃣ Créer sessions ---
        # On va créer une session par date
        sessions_dict = {}  # temporaire pour regrouper les exercices par date

        for ex_name, records in all_exercise_dict.items():
            ex_id = exercise_ids[ex_name]

            for record in records:
                date = record["Date"]
                rpe = record["Etat"] if "Etat" in record else ""
                notes = record["Notes"] if "Notes" in record else ""


                # Créer la session si elle n'existe pas encore
                if date not in sessions_dict:
                    sessions_dict[date] = {
                        "date": date,
                        "exercises": []
                    }

                # Créer les sets
                sets = []
                for k, v in record.items():
                    if k.startswith("S") and isinstance(v, (int, float)):
                        sets.append({"r": v, "w": record["Kg"], "tempo":"", "rest": ""})

                # Ajouter l'exercice à la session
                sessions_dict[date]["exercises"].append({
                    "exercise_id": ex_id,
                    "sets": sets,
                    "rpe": rpe,
                    "notes": notes
                })

        # Trier par date
        database["sessions"] = sorted(
            sessions_dict.values(),
            key=lambda x: x["date"]
        )

        self.export_database_to_csv(database, folder_path="my_training_data")

        return database
    
    def export_database_to_csv(self, database, folder_path):
        """
        Exporte la base de données en CSV dans un dossier.
        - sessions.csv
        - bodyweight.csv
        - exercise_library.csv

        :param database: dict contenant 'exercise_library', 'sessions', 'bodyweight_history'
        :param folder_path: dossier où stocker les fichiers CSV
        """
        # Créer le dossier s'il n'existe pas
        if not os.path.exists(folder_path):
            os.makedirs(folder_path)

        # -----------------------------
        # 1️⃣ Exporter exercise_library
        # -----------------------------
        lib_file = os.path.join(folder_path, "exercise_library.csv")
        with open(lib_file, "w", newline="", encoding="utf-8") as f:
            fieldnames = ["exercise_id", "name", "muscle_group", "primary_group", "secondary_group", "type"]
            writer = csv.DictWriter(f, fieldnames=fieldnames)
            writer.writeheader()

            for ex_id, ex_data in database["exercise_library"].items():
                writer.writerow({
                    "exercise_id": ex_id,
                    "name": ex_data.get("name",""),
                    "muscle_group": ex_data.get("muscle_group",""),
                    "primary_group": ",".join(ex_data.get("primary_group",[])),
                    "secondary_group": ",".join(ex_data.get("secondary_group",[])),
                    "type": ex_data.get("type","")
                })

        # -----------------------------
        # 2️⃣ Exporter sessions
        # -----------------------------
        sessions_file = os.path.join(folder_path, "sessions.csv")
        with open(sessions_file, "w", newline="", encoding="utf-8") as f:
            fieldnames = ["date", "exercise_name", "set_index", "reps", "weight", "tempo", "rest", "rpe", "notes"]
            writer = csv.DictWriter(f, fieldnames=fieldnames)
            writer.writeheader()

            for session in database["sessions"]:
                date_str = session["date"].strftime("%Y-%m-%d")
                for ex in session["exercises"]:
                    ex_name = database["exercise_library"][ex["exercise_id"]]["name"]
                    rpe = ex.get("rpe","")
                    notes = ex.get("notes","")
                    for i, s in enumerate(ex["sets"], start=1):
                        writer.writerow({
                            "date": date_str,
                            "exercise_name": ex_name,
                            "set_index": i,
                            "reps": s.get("r",""),
                            "weight": s.get("w",""),
                            "tempo": s.get("tempo",""),
                            "rest": s.get("rest",""),
                            "rpe": rpe,
                            "notes": notes
                        })

        # -----------------------------
        # 3️⃣ Exporter bodyweight_history
        # -----------------------------
        bodyweight_file = os.path.join(folder_path, "bodyweight.csv")
        with open(bodyweight_file, "w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=["date", "weight"])
            writer.writeheader()
            for rec in database.get("bodyweight_history", []):
                writer.writerow({
                    "date": rec["date"].strftime("%Y-%m-%d"),
                    "weight": rec["weight"]
                })

        print(f"Export terminé dans le dossier : {folder_path}")
    
    def import_database_from_csv(self,folder_path="database_export"):
        """
        Importe la base de données depuis les CSV d'un dossier et reconstruit la structure complète.

        :param folder_path: dossier contenant exercise_library.csv, sessions.csv, bodyweight.csv
        :return: database dict
        """
        database = {
            "exercise_library": {},
            "sessions": [],
            "bodyweight_history": []
        }

        # -----------------------------
        # 1️⃣ Charger exercise_library
        # -----------------------------
        lib_file = os.path.join(folder_path, "exercise_library.csv")
        with open(lib_file, newline="", encoding="utf-8") as f:
            reader = csv.DictReader(f)
            for row in reader:
                ex_id = int(row["exercise_id"])
                database["exercise_library"][ex_id] = {
                    "name": row.get("name",""),
                    "muscle_group": row.get("muscle_group",""),
                    "primary_group": row.get("primary_group","").split(",") if row.get("primary_group") else [],
                    "secondary_group": row.get("secondary_group","").split(",") if row.get("secondary_group") else [],
                    "type": row.get("type","")
                }

        # -----------------------------
        # 2️⃣ Charger sessions
        # -----------------------------
        sessions_file = os.path.join(folder_path, "sessions.csv")
        # On utilise un dict temporaire pour regrouper par date et par exercice
        sessions_dict = defaultdict(lambda: {"date": None, "exercises": defaultdict(lambda: {"exercise_id": None, "sets": [], "rpe": "", "notes": ""})})

        with open(sessions_file, newline="", encoding="utf-8") as f:
            reader = csv.DictReader(f)
            for row in reader:
                date = datetime.strptime(row["date"], "%Y-%m-%d")
                ex_name = row["exercise_name"]
                # Trouver l'exercise_id
                ex_id = next((eid for eid, ed in database["exercise_library"].items() if ed["name"] == ex_name), None)
                if ex_id is None:
                    continue  # skip si exercice absent dans library

                # Créer session
                session = sessions_dict[date]
                session["date"] = date
                ex_data = session["exercises"][ex_id]
                ex_data["exercise_id"] = ex_id
                ex_data["rpe"] = row.get("rpe","")
                ex_data["notes"] = row.get("notes","")
                # Ajouter set
                ex_data["sets"].append({
                    "r": float(row.get("reps",0)),
                    "w": float(row.get("weight",0)),
                    "tempo": row.get("tempo",""),
                    "rest": row.get("rest","")
                })

        # Transformer en liste triée
        for date in sorted(sessions_dict.keys()):
            session = sessions_dict[date]
            # Convertir exercises dict en liste
            session["exercises"] = list(session["exercises"].values())
            database["sessions"].append(session)

        # -----------------------------
        # 3️⃣ Charger bodyweight_history
        # -----------------------------
        bodyweight_file = os.path.join(folder_path, "bodyweight.csv")
        if os.path.exists(bodyweight_file):
            with open(bodyweight_file, newline="", encoding="utf-8") as f:
                reader = csv.DictReader(f)
                for row in reader:
                    database["bodyweight_history"].append({
                        "date": datetime.strptime(row["date"], "%Y-%m-%d"),
                        "weight": float(row["weight"])
                    })

        print(f"Import terminé depuis le dossier : {folder_path}")
        return database

    # Gestion de la base de données
    def get_exercise_names(self, database):
        return [ex["name"] for ex in database["exercise_library"].values()]