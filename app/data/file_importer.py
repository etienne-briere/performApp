
# KivyMD et Kivy
from kivymd.toast import toast
from kivy.app import App
from kivy.event import EventDispatcher
from kivy.properties import ListProperty
from kivy.utils import platform

# Gestionnaires
from app.data.data_converter import DataConverter

# Utilitaires
from plyer import filechooser
from openpyxl import load_workbook
import os

# Android-specific imports
if platform == "android":
    from jnius import autoclass, cast
    from android import activity


class FileImporter(EventDispatcher):

    exercise_names = ListProperty([])
    
    def __init__(self, repo, **kwargs):
        super().__init__(**kwargs)

        # Gestionnaire de données
        self.repo = repo


    def read_excel(self,file_path) -> dict:
        """Lit un .xlsx et retourne un dict brut {sheet_name: [row_dict]}."""
        wb = load_workbook(file_path, data_only=True)
        # all_sheets: dict[str, list[dict[str, str | float | datetime | None]]] = {}
        all_sheets = {}

        for sheetname in wb.sheetnames:
            ws = wb[sheetname]

            # Récupérer les en-têtes (première ligne)
            headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]

            # Récupérer les lignes suivantes sous forme de dict
            rows = []
            for row in ws.iter_rows(min_row=2, values_only=True):
                row_dict = dict(zip(headers, row))
                rows.append(row_dict)

            all_sheets[sheetname] = rows

        return all_sheets

    def open(self):
        """Lance le sélecteur natif selon la plateforme."""
        if platform == "android":
            self._open_file_android(self.select_file)
        else:
            filechooser.open_file(on_selection=self.select_file)

    def _open_file_android(self, on_selection):
        PythonActivity = autoclass('org.kivy.android.PythonActivity')
        Intent = autoclass('android.content.Intent')

        intent = Intent(Intent.ACTION_GET_CONTENT)
        intent.setType(
            "*/*")  # tu peux restreindre à "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        intent.addCategory(Intent.CATEGORY_OPENABLE)

        # Stocke ton callback dans une variable globale
        global ANDROID_FILE_CALLBACK
        ANDROID_FILE_CALLBACK = on_selection

        # Attache un listener sur le résultat
        activity.bind(on_activity_result=self._on_activity_result)

        # Lance le chooser
        PythonActivity.mActivity.startActivityForResult(intent, 1001)

    def _on_activity_result(self, requestCode, resultCode, data):
        from android import activity
        if requestCode != 1001:
            return

        if resultCode == -1:  # RESULT_OK
            uri = data.getData().toString()
            ANDROID_FILE_CALLBACK([uri])
        else:
            ANDROID_FILE_CALLBACK([])

    def copy_uri_to_temp_file(self, uri_string):
        PythonActivity = autoclass('org.kivy.android.PythonActivity')
        context = PythonActivity.mActivity
        Uri = autoclass('android.net.Uri')
        uri = Uri.parse(uri_string)

        content_resolver = context.getContentResolver()
        input_stream = content_resolver.openInputStream(uri)

        temp_dir = context.getCacheDir().getAbsolutePath()
        temp_path = os.path.join(temp_dir, "imported_excel.xlsx")

        FileOutputStream = autoclass('java.io.FileOutputStream')
        output_stream = FileOutputStream(temp_path)

        buffer = bytearray(1024)
        read = input_stream.read(buffer)
        while read != -1:
            output_stream.write(buffer, 0, read)
            read = input_stream.read(buffer)

        input_stream.close()
        output_stream.close()

        return temp_path

    def select_file(self, selection):
        """Callback quand un fichier est sélectionné"""

        print(f"Sélection : {selection}")

        # selection est soit None, soit une liste vide, soit [chemin]
        if not selection:
            toast("Aucun fichier sélectionné")
            print("Aucun fichier sélectionné (ou annulation).")
            return

        file_path = selection[0]  # premier fichier choisi
        print("Fichier choisi :", file_path)
        if not file_path:
            toast("Fichier invalide")
            print("Fichier invalide")
            return

        # Si on est sur Android → convertir content:// en vrai fichier
        if file_path.startswith("content://"):
            print("URI détectée → copie vers fichier temporaire")
            file_path = self.copy_uri_to_temp_file(file_path)

        print("Chemin final utilisé :", file_path)

        # Maintenant safe : file_path est une string
        if any(file_path.endswith(ext) for ext in [".xls", ".xlsx", ".csv"]):
            # Charger le fichier
            self.all_exercise_dict = self.read_excel(file_path)
            print(f"Dictionnaire complet importé : {self.all_exercise_dict}")

            convert = DataConverter()
            database = convert.from_legacy(self.all_exercise_dict)
            print(f"Dictionnaire converti : {database}")

            # Alerter l'utilisateur
            toast("Fichier chargé avec succès !")

            # Charger la base complète depuis le dossier CSV
            # database = self.import_database_from_csv(folder_path="my_training_data")
            
            # liste des noms des exercices
            self.exercise_names = self.repo.get_exercise_names(database)

            # Informer l'app que l'exercise sélectionné par défaut est le 1er
            if self.exercise_names:
                app = App.get_running_app()
                app.selected_exercise = self.exercise_names[0]    

        else:
            toast(text="Extension non supportée")
    